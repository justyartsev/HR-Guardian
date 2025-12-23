import os
import io
import re
from pathlib import Path
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
import markdown
import fitz  # PyMuPDF
from PIL import Image
import pytesseract


class DocumentReader:
    def __init__(self, tesseract_cmd: str = None):
        """Инициализирует читатель документов (параметры: tesseract_cmd; возвращает: None)."""
        if tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    def read(self, file_path: str) -> str:
        """Читает документ (TXT, DOCX, MD, PDF) по расширению файла (параметры: file_path; возвращает: str текст)."""
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError(f"Файл не найден: {file_path}")

        ext = path.suffix.lower()
        if ext == '.txt':
            return self._read_txt(file_path)
        elif ext == '.docx':
            return self._read_docx(file_path)
        elif ext == '.md':
            return self._read_md(file_path)
        elif ext == '.pdf':
            return self._read_pdf(file_path)
        else:
            raise ValueError(f"Неподдерживаемый формат файла: {ext}")

    def _read_txt(self, file_path: str) -> str:
        """Читает TXT файл с UTF-8 кодировкой (параметры: file_path; возвращает: str)."""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            print(f"Error reading .txt file {file_path}: {e}")
            return ""

    def _read_docx(self, file_path: str) -> str:
        """Читает DOCX: параграфы и таблицы с markdown-форматированием (параметры: file_path; возвращает: str)."""
        try:
            doc = Document(file_path)
            full_text = []

            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)

            for table in doc.tables:
                # Преобразуем таблицу в markdown-формат
                table_text = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        table_text.append("| " + " | ".join(cells) + " |")
                if table_text:
                    full_text.append("\nТаблица:\n" + "\n".join(table_text) + "\n")

            return '\n'.join(full_text)
        except Exception as e:
            print(f"Error reading .docx: {e}")
            return ""


    def _read_md(self, file_path: str) -> str:
        """Читает Markdown: конвертирует в HTML и удаляет теги (параметры: file_path; возвращает: str)."""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                md_text = f.read()
            html = markdown.markdown(md_text)
            # Удаляем HTML-теги регулярным выражением
            clean_text = re.sub(r'<[^>]+>', '', html)
            return clean_text
        except Exception as e:
            print(f"Error reading .md file {file_path}: {e}")
            return ""

    def _read_pdf(self, file_path: str) -> str:
        """Читает PDF: текст со страниц или OCR если требуется (параметры: file_path; возвращает: str)."""
        try:
            doc = fitz.open(file_path)
            full_text = ""
            for page in doc:
                # Пытаемся извлечь текст
                text = page.get_text().strip()
                if text:
                    full_text += text + "\n"
                else:
                    # Если текста нет, применяем OCR к изображению страницы
                    pix = page.get_pixmap(dpi=150)
                    img = Image.open(io.BytesIO(pix.tobytes("png")))
                    ocr_text = pytesseract.image_to_string(img, lang='rus+eng')
                    full_text += ocr_text + "\n"
            return full_text
        except Exception as e:
            print(f"Error reading .pdf: {e}")
            return ""